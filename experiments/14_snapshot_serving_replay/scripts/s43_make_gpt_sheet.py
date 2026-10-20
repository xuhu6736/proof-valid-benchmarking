#!/usr/bin/env python3
"""s43_make_gpt_sheet.py -- annotation sheet for the GPT regeneration round.

Goal: upgrade the second-generator (gpt-5-mini) archive to human-labelled
truth with the smallest honest annotation effort.

Label-reuse rule (agreed with the authors):
  * GPT answers byte-identical to the kimi answer after normalisation
    (whitespace/case/commas) inherit that key's adjudicated label -- the
    annotators already judged exactly this text against this evidence
    (104 of the 174 Table-B keys, 22 of the 51 s9 keys).
  * Every differing answer must be judged anew.  This sheet collects
    exactly those keys:
      - 69 merged-store keys (66 same-target-value, 3 different-value)
      - 29 endpoint (s9) keys with differing text
      -  1 merged-store key that kimi failed acceptance on but GPT
        answered (newly served in the GPT variant; never labelled)
    Total: 99 entries.

Protocol is identical to the earlier rounds: same R1 operational
definitions, same bilingual layout, same sqrt/box marking, gate decisions
and the kimi answers/labels withheld.

Outputs:
  annotation/PVSB_gpt_sheet.docx               (working copy)
  <paper dir>/PVSB_GPT再生标注表_99条.docx      (deliverable copy)
  annotation/gpt_sheet_keys.json               (key list for parsing)
"""
import json, os, re, html, shutil
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.normpath(os.path.join(HERE, '..', 'data'))
ANN = os.path.normpath(os.path.join(HERE, '..', 'annotation'))
OUT = os.path.join(ANN, 'PVSB_gpt_sheet.docx')
PAPER_DIR = os.path.normpath(os.path.join(HERE, '..', '..', '..', '..'))
T = '20260801'
TICK_LABEL = {'20260501': '2026-05-01', '20260701': '2026-07-01',
              '20260801': '2026-08-01'}


def norm(s):
    return re.sub(r'[,\s]+', '', (s or '').lower())


def clean(wt):
    wt = html.unescape(wt)
    wt = re.sub(r"<ref[^>]*>.*?</ref>", " ", wt, flags=re.S)
    wt = re.sub(r"<ref[^>]*/>", " ", wt)
    wt = re.sub(r"\{\{[^{}]*\}\}", " ", wt)
    wt = re.sub(r"\{\{[^{}]*\}\}", " ", wt)
    wt = re.sub(r"\[\[([^|\]]*\|)?([^\]]*)\]\]", r"\2", wt)
    wt = re.sub(r"\[https?://[^\s\]]+\s*([^\]]*)\]", r"\1", wt)
    wt = re.sub(r"'{2,}", "", wt)
    wt = re.sub(r"<[^>]+>", " ", wt)
    return re.sub(r"\s+", " ", wt).strip()


def main():
    ktab = json.load(open(f'{DATA}/tableb_regeneration.json'))
    gtab = json.load(open(f'{DATA}/tableb_regeneration_gpt.json'))
    k9 = json.load(open(f'{DATA}/regeneration.json'))
    g9 = json.load(open(f'{DATA}/regeneration_gpt.json'))
    tl = json.load(open(f'{DATA}/tableb_labels.json'))      # 174 human labels
    emap = json.load(open(f'{DATA}/tableb_endpoint_map.json'))
    ev = {x['pid']: x for x in json.load(open(f'{DATA}/evidence.json'))}
    gen = json.load(open(f'{DATA}/generation.json'))
    tr = json.load(open(f'{DATA}/translations.json'))

    # keys actually delivered at some arm's endpoint (Table B scope only;
    # gap-fill-only keys serve the count-based fault sweeps and carry no
    # human labels by design)
    used = set()
    for arm, em in emap.items():
        for pid, rec in em.items():
            if rec['source'] == 'template':
                used.add(f"{pid}|{rec['origin']}")

    # ---------------- collect keys needing fresh labels -------------------
    entries = []          # (sort_key, key_string, question_ctx record)
    for k, g in gtab.items():
        if k not in tl and k not in used:
            continue                     # not a Table B endpoint key
        if not g['regen_ok']:
            continue                     # accept-fail / no-evidence: abstain
        kr = ktab.get(k, {})
        if kr.get('regen_ok') and norm(kr.get('regen_answer')) == \
                norm(g['regen_answer']):
            continue                     # identical text: label inherited
        entries.append((int(g['pid']), g['tick'], k, {
            'tick': g['tick'], 'answer': g['regen_answer'],
            'values': g['target_values_text'],
            'sents': g['target_sentences'],
            'template': g.get('template_answer')}))
    for pid, g in g9.items():
        if not g['regen_ok']:
            continue
        kr = k9.get(pid, {})
        if kr.get('regen_ok') and norm(kr.get('regen_answer')) == \
                norm(g['regen_answer']):
            continue
        entries.append((int(pid), T, f'{pid}|{T}', {
            'tick': T, 'answer': g['regen_answer'],
            'values': g['target_values_text'],
            'sents': g['target_sentences'], 'template': None}))
    entries.sort(key=lambda e: (e[0], e[1]))
    print('entries needing fresh labels:', len(entries))

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(10)

    doc.add_heading(
        'PVSB GPT 再生标注表 Second-Generator (GPT) Regeneration Sheet', 0)
    doc.add_paragraph(
        f'说明 Instructions：共 {len(entries)} 条。这是第二生成模型（GPT）按'
        '与首轮完全相同的提示词与当期证据再生的答案；凡与首轮答案文本一致'
        '的条目已直接沿用既有标注，本表仅收录文本有差异的条目。'
        '请独立判断每个再生答案在其证据期快照下的正确性，'
        '在对应选项前把 □ 改为 √。首轮答案、门控判定与各臂机制标签均未向您'
        '展示。 For each entry, mark exactly one option by changing the '
        'leading □ to √.  First-round answers, gate decisions, and per-arm '
        'mechanism labels are withheld.')
    leg = doc.add_table(rows=5, cols=2)
    leg.style = 'Table Grid'
    leg.rows[0].cells[0].text = 'R1 档位 Level'
    leg.rows[0].cells[1].text = '操作定义 Operational definition'
    for i, (lv, df) in enumerate([
            ('正确 correct', '答案与当期目标值一致且完整回答了问题'),
            ('错误 incorrect', '答案与当期目标值矛盾，或给出目标值之外的值'),
            ('不完整 incomplete', '方向正确但漏掉了部分当期目标值'),
            ('无法判断 indeterminate', '证据不足以按上述三档判定')]):
        leg.rows[i + 1].cells[0].text = lv
        leg.rows[i + 1].cells[1].text = df
    doc.add_paragraph('')

    keys_out = []
    for _, _, k, r in entries:
        pid, tick = k.split('|')
        pid = int(pid)
        x = ev[pid]
        keys_out.append(k)
        t = doc.add_table(rows=8, cols=2)
        t.style = 'Table Grid'
        t.rows[0].cells[0].text = (
            f'条目 Entry #{pid} | 证据期 {TICK_LABEL.get(tick, tick)}   '
            f'{x["label"]} (Q{x["qid"]})   {x["prop"]} {x["prop_name"]}')
        t.rows[0].cells[1].text = ''
        zh = tr.get(str(pid), {}).get('q_zh', '')
        t.rows[1].cells[0].text = '问题 Question'
        t.rows[1].cells[1].text = x['question'] + ('\n' + zh if zh else '')
        t.rows[2].cells[0].text = '改写查询 Paraphrase'
        t.rows[2].cells[1].text = gen[str(pid)]['paraphrase']
        t.rows[3].cells[0].text = '缓存答案（基期） Cached answer (base)'
        t.rows[3].cells[1].text = gen[str(pid)]['answer']
        t.rows[4].cells[0].text = '当期目标值 Target values @ tick'
        t.rows[4].cells[1].text = '; '.join(r['values']) or '—'
        t.rows[5].cells[0].text = '当期证据句 Evidence sentences @ tick'
        t.rows[5].cells[1].text = '\n'.join(
            f'[{i+1}] {clean(s)}' for i, s in enumerate(r['sents'])) or '—'
        t.rows[6].cells[0].text = '再生答案 Regenerated answer (GPT)'
        t.rows[6].cells[1].text = r['answer']
        t.rows[7].cells[0].text = '标注 Annotation'
        t.rows[7].cells[1].text = (
            'R1 再生答案当期正确性 correctness at origin tick:  '
            '□ 正确 correct   □ 错误 incorrect   □ 不完整 incomplete   '
            '□ 无法判断 indeterminate\n'
            '备注 notes: ')
        doc.add_paragraph('')

    doc.save(OUT)
    json.dump(keys_out, open(os.path.join(ANN, 'gpt_sheet_keys.json'), 'w'),
              ensure_ascii=False, indent=1)
    print('wrote', OUT, 'with', len(entries), 'entries')
    dst = os.path.join(PAPER_DIR, f'PVSB_GPT再生标注表_{len(entries)}条.docx')
    shutil.copyfile(OUT, dst)
    print('deliverable copy:', dst)


if __name__ == '__main__':
    main()
