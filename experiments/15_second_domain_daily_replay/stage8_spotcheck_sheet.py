#!/usr/bin/env python3
"""Stage 8 (A-2 final): double-annotation sheets for the 20-item human
spot-check of fallback regenerations (seed-42 sample, identical to the
pilot sheet), plus a machine-readable master ledger for parse-back.

Differences from the pilot sheet (second_domain_pilot/...docx):
  * the auto-judgment column is WITHHELD from annotators (no anchoring);
  * four-level checkbox annotation (correct / incomplete / incorrect /
    indeterminate), consistent with the 51/174/99-item sheets;
  * the question stem and per-part target values are shown explicitly.

Outputs (workspace root):
  PVSB_主场景回退标注表_20条（人工1）.docx
  PVSB_主场景回退标注表_20条（人工2）.docx
and (here) fallback_spotcheck_sample20.json
"""
import json
import os
import random

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))


def build_sample():
    data = json.load(open(os.path.join(HERE, "fallback_regen.json")))
    rng = random.Random(42)
    rows = []
    for gen in ("kimi", "gpt"):
        for k, v in data.get(gen, {}).items():
            if "correct" in v:
                rows.append((gen, k, v))
    sample = rng.sample(rows, min(20, len(rows)))

    panel = json.load(open(os.path.join(HERE, "panel21.json")))["panel"]
    questions = {}
    for gen in ("kimi", "gpt"):
        ans = json.load(open(os.path.join(HERE, f"answers_{gen}.json")))[
            "answers"]
        for k, a in ans.items():
            questions[(gen, k)] = a["question"].split("Evidence from")[0
                ].strip()

    items = []
    for n, (gen, key, v) in enumerate(sample, 1):
        idx = int(key.rsplit("#", 1)[1])
        claim_keys = [c["key"] for c in panel[idx]["claims"]]
        items.append({
            "n": n, "gen": gen, "key": key, "page": panel[idx]["page"],
            "tick": v["tick"], "claim_keys": claim_keys,
            "question": questions[(gen, key)],
            "current_values": v["current_values"],
            "regen": v["regen"],
            "auto_correct": v["correct"], "auto_missing": v["missing"],
        })
    return items


LEVELS = [
    ("正确 correct",
     "答案完整陈述了当期全部目标值（措辞不必一致；数字的等价写法视为一致）。"
     " The answer states every target value at the rejection tick "
     "(wording need not match)."),
    ("不完整 incomplete",
     "方向正确但漏掉部分当期目标值，或对某项明示“证据未述”而当期实际有值。"
     " Partly right but omits some target values, or claims 'not stated' "
     "where a value exists."),
    ("错误 incorrect",
     "至少一项与当期目标值矛盾，或给出目标值之外的值。"
     " At least one part contradicts the target value or gives an "
     "off-target value."),
    ("无法判断 indeterminate",
     "信息不足以按上述三档判定。 Evidence insufficient for the three "
     "levels above."),
]


def make_doc(items, annotator, path):
    doc = Document()
    doc.add_heading(
        f"PVSB 主场景回退再生成标注表 Fallback-Regeneration Annotation "
        f"Sheet（20 条，{annotator}）", level=1)
    doc.add_paragraph(
        "说明 Instructions：共 20 条，均为日粒度主场景中断言门控首次拒答后、"
        "以当期导语为证据重新生成的答案。请独立判断每个再生答案在其拒绝日"
        "（rejection tick）的正确性：对照“当期目标值”，在对应选项前把 □ "
        "改为 √，每条限选一项。门控判定、自动判正结果与缓存答案标签均未向"
        "您展示。 For each entry, judge the regenerated answer against the "
        "target values at its rejection tick and mark exactly one option by "
        "changing the leading □ to √.  Gate decisions, auto-judgments and "
        "cached-answer labels are withheld.")
    t0 = doc.add_table(rows=1, cols=2)
    t0.style = "Light Grid Accent 1"
    t0.rows[0].cells[0].text = "档位 Level"
    t0.rows[0].cells[1].text = "操作定义 Operational definition"
    for name, desc in LEVELS:
        c = t0.add_row().cells
        c[0].text = name
        c[1].text = desc
    doc.add_paragraph("")

    for it in items:
        t = doc.add_table(rows=5, cols=2)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].merge(hdr[1])
        t.rows[0].cells[0].text = (
            f"条目 Entry #{it['n']} | 生成器 {it['gen']} | "
            f"页面 {it['page']} | 拒绝日 rejection tick {it['tick']}")
        rows = t.rows
        rows[1].cells[0].text = "问题 Question"
        rows[1].cells[1].text = it["question"]
        vals = []
        for j, (ck, cv) in enumerate(
                zip(it["claim_keys"], it["current_values"]), 1):
            vals.append(f"({j}) [{ck}] " + (
                str(cv) if cv is not None else
                "（当期信息框无此键值；答“未列出”视为正确 / no infobox "
                "value at tick）"))
        rows[2].cells[0].text = "当期目标值 Target values @ tick"
        rows[2].cells[1].text = "\n".join(vals)
        rows[3].cells[0].text = "再生答案 Regenerated answer"
        rows[3].cells[1].text = it["regen"]
        rows[4].cells[0].text = "标注 Annotation"
        rows[4].cells[1].text = ("再生答案当期正确性 correctness at "
                                 "rejection tick:  " + "   ".join(
                                     f"□ {name}" for name, _ in LEVELS))
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
        doc.add_paragraph("")
    doc.save(path)


def main():
    items = build_sample()
    ledger = os.path.join(HERE, "fallback_spotcheck_sample20.json")
    json.dump({"sample_seed": 42, "items": items},
              open(ledger, "w"), ensure_ascii=False, indent=1)
    print("saved", ledger)
    for tag in ("人工1", "人工2"):
        path = os.path.join(
            ROOT, f"PVSB_主场景回退标注表_20条（{tag}）.docx")
        make_doc(items, tag, path)
        print("saved", path)
    ok = sum(i["auto_correct"] for i in items)
    print(f"auto-judged correct: {ok}/{len(items)}")


if __name__ == "__main__":
    main()
