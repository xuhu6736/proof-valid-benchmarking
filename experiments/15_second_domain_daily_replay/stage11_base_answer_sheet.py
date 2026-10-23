#!/usr/bin/env python3
"""Stage 11 (operating-window closure, part 2): double-annotation sheets
for a sample of ACCEPTED BASE (cached) answers, to estimate b -- the
error rate of non-stale cached answers -- under the same protocol as the
fallback spot check.  No model calls.

Sample: 40 accepted base answers (20 per generator, seed 42), stratified
to over-include per-part-flagged answers (10 flagged + 10 unflagged per
generator) so the estimate resolves the cross-part-masking failure mode
with precision while remaining unbiased via design weights recorded in
the ledger.

Outputs (workspace root):
  PVSB_主场景基缓存答案标注表_40条（人工1）.docx
  PVSB_主场景基缓存答案标注表_40条（人工2）.docx
and (here) base_answer_sample40.json (master ledger incl. withheld
per-part verdicts and stratum weights).
"""
import json
import os
import random

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))

LEVELS = [
    ("正确 correct",
     "答案完整陈述了基期全部目标值（措辞不必一致；数字的等价写法视为一致）。"
     " The answer states every t0 target value (wording need not match)."),
    ("不完整 incomplete",
     "方向正确但漏掉部分基期目标值，或对某项明示“证据未述”而基期实际有值。"
     " Partly right but omits some t0 target values, or claims 'not "
     "stated' where a value exists."),
    ("错误 incorrect",
     "至少一项与基期目标值矛盾，或给出目标值之外的值。"
     " At least one part contradicts the t0 target value or gives an "
     "off-target value."),
    ("无法判断 indeterminate",
     "信息不足以按上述三档判定。 Evidence insufficient for the three "
     "levels above."),
]


def build_sample():
    panel = json.load(open(os.path.join(HERE, "panel21.json")))["panel"]
    rescore = json.load(open(os.path.join(HERE, "perpart_rescore.json")))
    rng = random.Random(42)
    items = []
    n_per_stratum = 10
    for gen in ("kimi", "gpt"):
        ans = json.load(open(os.path.join(
            HERE, f"answers_{gen}.json")))["answers"]
        flagged, clean = [], []
        for key, a in ans.items():
            if not a.get("accepted"):
                continue
            rec = rescore["base_answers"][f"{gen}|{key}"]
            (flagged if not rec["per_part_accept"] else clean).append(
                (key, a))
        pick = rng.sample(flagged, n_per_stratum) + \
            rng.sample(clean, n_per_stratum)
        rng.shuffle(pick)
        # design weights: stratum size / sampled
        w = {"flagged": len(flagged) / n_per_stratum,
             "clean": len(clean) / n_per_stratum}
        for key, a in pick:
            idx = int(key.rsplit("#", 1)[1])
            claims = panel[idx]["claims"]
            rec = rescore["base_answers"][f"{gen}|{key}"]
            items.append({
                "gen": gen, "key": key, "page": panel[idx]["page"],
                "stratum": "flagged" if not rec["per_part_accept"]
                else "clean",
                "weight": w["flagged"] if not rec["per_part_accept"]
                else w["clean"],
                "claim_keys": [c["key"] for c in claims],
                "t0_values": [c["value_t0"] for c in claims],
                "question": a["question"].split("Evidence from")[0].strip(),
                "answer": a["answer"],
                "per_part_accept": rec["per_part_accept"],
                "per_part_missing": rec["per_part_missing"],
            })
    for n, it in enumerate(items, 1):
        it["n"] = n
    return items


def make_doc(items, annotator, path):
    doc = Document()
    doc.add_heading(
        f"PVSB 主场景基缓存答案标注表 Base Cached-Answer Annotation "
        f"Sheet（40 条，{annotator}）", level=1)
    doc.add_paragraph(
        "说明 Instructions：共 40 条，均为日粒度主场景中通过词汇验收"
        "入库的基期缓存答案（t0 = 2026-07-01）。请独立判断每条缓存答案"
        "相对基期目标值的正确性，在对应选项前把 □ 改为 √，每条限选一"
        "项。自动判定与门控标签均未向您展示。 For each entry, judge the "
        "cached base answer against its t0 target values and mark exactly "
        "one option by changing the leading □ to √.  Automatic verdicts "
        "and gate labels are withheld.")
    t0 = doc.add_table(rows=1, cols=2)
    t0.style = "Light Grid Accent 1"
    t0.rows[0].cells[0].text = "档位 Level"
    t0.rows[0].cells[1].text = "操作定义 Operational definition"
    for name, desc in LEVELS:
        c = t0.add_row().cells
        c[0].text = name
        c[1].text = desc
    doc.add_paragraph("")

    for it in items:
        t = doc.add_table(rows=5, cols=2)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].merge(hdr[1])
        t.rows[0].cells[0].text = (
            f"条目 Entry #{it['n']} | 生成器 {it['gen']} | "
            f"页面 {it['page']} | 基期 t0 2026-07-01")
        rows = t.rows
        rows[1].cells[0].text = "问题 Question"
        rows[1].cells[1].text = it["question"]
        vals = []
        for j, (ck, cv) in enumerate(
                zip(it["claim_keys"], it["t0_values"]), 1):
            vals.append(f"({j}) [{ck}] " + (
                str(cv) if cv is not None else
                "（基期信息框无此键值；答“未列出”视为正确 / no infobox "
                "value at t0）"))
        rows[2].cells[0].text = "基期目标值 Target values @ t0"
        rows[2].cells[1].text = "\n".join(vals)
        rows[3].cells[0].text = "缓存答案 Cached base answer"
        rows[3].cells[1].text = it["answer"]
        rows[4].cells[0].text = "标注 Annotation"
        rows[4].cells[1].text = ("缓存答案基期正确性 correctness at t0:  "
                                 + "   ".join(f"□ {name}"
                                             for name, _ in LEVELS))
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
        doc.add_paragraph("")
    doc.save(path)


def main():
    items = build_sample()
    ledger = os.path.join(HERE, "base_answer_sample40.json")
    json.dump({"sample_seed": 42,
               "design": "10 per-part-flagged + 10 clean per generator; "
                         "design weights in 'weight'",
               "items": items},
              open(ledger, "w"), ensure_ascii=False, indent=1)
    print("saved", ledger)
    for tag in ("人工1", "人工2"):
        path = os.path.join(
            ROOT, f"PVSB_主场景基缓存答案标注表_40条（{tag}）.docx")
        make_doc(items, tag, path)
        print("saved", path)
    from collections import Counter
    print(Counter((i["gen"], i["stratum"]) for i in items))


if __name__ == "__main__":
    main()
