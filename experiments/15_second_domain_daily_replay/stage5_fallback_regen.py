#!/usr/bin/env python3
"""Stage 5 (A-2): real fallback regeneration for the first rejection of
each cached answer under the sent-asserted gate, judged against the
deterministic tick-k truth; plus a human spot-check sheet (docx).

Regeneration prompt = the answer's original question + the CURRENT
tick-k lead as evidence. Auto-judge: the regenerated text must mention
each claim's current infobox value (normalized). Output:
fallback_regen.json + PVSB_主场景回退抽检表.docx (20 sampled rows).
"""
import json
import os
import re
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
TICKS = [f"2026-07-{d:02d}" for d in range(1, 22)]

import urllib.request

PROXY = os.environ.get("GEN_PROXY", "http://127.0.0.1:7890")
OPENER_PROXY = urllib.request.build_opener(
    urllib.request.ProxyHandler({"http": PROXY, "https": PROXY}))
OPENER_DIRECT = urllib.request.build_opener()


def norm(s):
    return re.sub(r"[,\s]+", "", s.lower())


def chat_gen(gen, prompt, system, max_tokens=200, tries=5):
    if gen == "kimi":
        url = os.environ["KIMI_BASE_URL"].rstrip("/") + "/messages"
        headers = {"x-api-key": os.environ["KIMI_API_KEY"],
                   "anthropic-version": "2023-06-01",
                   "content-type": "application/json"}
        payload = {"model": "kimi-k2-turbo-preview",
                   "max_tokens": max_tokens,
                   "thinking": {"type": "disabled"}, "system": system,
                   "messages": [{"role": "user", "content": prompt}]}
    else:
        url = os.environ.get("POE_BASE_URL",
                             "https://api.poe.com/v1").rstrip("/") \
            + "/chat/completions"
        headers = {"Authorization": "Bearer " + os.environ["POE_API_KEY"],
                   "content-type": "application/json"}
        payload = {"model": os.environ.get("GPT_MODEL", "gpt-5-mini"),
                   "max_tokens": max_tokens,
                   "messages": [{"role": "system", "content": system},
                                {"role": "user", "content": prompt}]}
    for a in range(tries):
        try:
            req = urllib.request.Request(url, data=json.dumps(payload).encode(),
                                         headers=headers, method="POST")
            opener = OPENER_DIRECT if gen == "kimi" else OPENER_PROXY
            with opener.open(req, timeout=90) as r:
                d = json.loads(r.read())
            if gen == "kimi":
                return "".join(b.get("text", "") for b in d.get("content", [])
                               if b.get("type") == "text").strip()
            return d["choices"][0]["message"]["content"].strip()
        except Exception:
            time.sleep(min(45, 3 + a * 5))
    return None



def current_value(series_page, key, day):
    snap = series_page["series"].get(day)
    if snap is None:
        return None
    return snap["infobox"].get(key)


def value_ok(answer, values):
    na = norm(answer)
    missing = []
    for v in values:
        if v is None:
            continue
        dv = re.findall(r"\d+(?:\.\d+)?", v.replace(",", ""))
        if dv:
            if not all(t in na for t in dv):
                missing.append(v)
        elif norm(v) not in na:
            missing.append(v)
    return (not missing), missing


SYSTEM = ("You are a precise assistant. Use only the provided evidence. "
          "State the specific values asked for. Keep the whole answer "
          "under 90 words.")


def first_rejection_tick(ans):
    for day in TICKS[1:]:
        alive = [c["alive"].get(day) for c in ans["claims"]]
        if any(x is False for x in alive):
            return day
    return None


def main():
    panel = json.load(open(os.path.join(HERE, "panel21.json")))["panel"]
    series = json.load(open(os.path.join(HERE, "series21.json")))["pages"]
    out_path = os.path.join(HERE, "fallback_regen.json")
    data = json.load(open(out_path)) if os.path.exists(out_path) else {}

    for gen in ("kimi", "gpt"):
        answers = json.load(open(os.path.join(HERE, f"answers_{gen}.json")))["answers"]
        g = data.setdefault(gen, {})
        for i, ans in enumerate(panel):
            key = f"{ans['page']}#{i}"
            if not answers.get(key, {}).get("accepted"):
                continue
            if key in g:
                continue
            day = first_rejection_tick(ans)
            if day is None:
                g[key] = {"never_rejected": True}
                continue
            page = series[ans["page"]]
            lead = page["series"].get(day, {}).get("lead", "")[:3500]
            q = answers[key]["question"].split("Evidence from")[0]
            prompt = (q + f"\nEvidence from the English Wikipedia article "
                      f"\"{ans['page']}\" (current revision at {day}):\n"
                      f"{lead}\n\nAnswer all three parts concisely, "
                      f"numbered (1) (2) (3).")
            text = chat_gen(gen, prompt, SYSTEM)
            if text is None:
                print(f"[fail] {gen} {key}", flush=True)
                continue
            cur_vals = [current_value(page, c["key"], day)
                        for c in ans["claims"]]
            ok, missing = value_ok(text, cur_vals)
            g[key] = {"tick": day, "regen": text, "current_values": cur_vals,
                      "correct": ok, "missing": missing}
            json.dump(data, open(out_path, "w"), ensure_ascii=False)
            print(f"[{gen}] {key} @{day} correct={ok}", flush=True)
            time.sleep(0.4)
        json.dump(data, open(out_path, "w"), ensure_ascii=False)

    # summary
    print("\n=== fallback regeneration (first asserted-rejection) ===")
    tot = {}
    for gen in ("kimi", "gpt"):
        g = data.get(gen, {})
        done = [v for v in g.values() if "correct" in v]
        ok = sum(v["correct"] for v in done)
        never = sum(1 for v in g.values() if v.get("never_rejected"))
        tot[gen] = {"regenerated": len(done), "correct": ok, "never": never}
        print(f"{gen}: {ok}/{len(done)} correct "
              f"({100*ok/max(1,len(done)):.1f}%); never rejected: {never}")
    json.dump(data, open(out_path, "w"), ensure_ascii=False)

    # ---- human spot-check sheet (20 rows sampled across generators) ----
    import random
    from docx import Document
    from docx.shared import Pt
    rng = random.Random(42)
    rows = []
    for gen in ("kimi", "gpt"):
        for k, v in data.get(gen, {}).items():
            if "correct" in v:
                rows.append((gen, k, v))
    sample = rng.sample(rows, min(20, len(rows)))
    doc = Document()
    doc.add_heading("PVSB 主场景回退再生成抽检表（20 条）", level=1)
    doc.add_paragraph("判定口径：再生成答案是否正确陈述了该 tick 的当前三值"
                      "（不要求措辞一致）。自动判定已给出，请在“人工判定”"
                      "列填写 正确/部分正确/错误。")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for j, h in enumerate(["编号", "生成器/条目/tick", "当前真值（三值）",
                           "再生成答案", "自动判定", "人工判定"]):
        hdr[j].text = h
    for n, (gen, k, v) in enumerate(sample, 1):
        cells = table.add_row().cells
        cells[0].text = str(n)
        cells[1].text = f"{gen} | {k} | {v['tick']}"
        cells[2].text = " | ".join(str(x) for x in v["current_values"])
        cells[3].text = v["regen"][:600]
        cells[4].text = "正确" if v["correct"] else f"缺 {v['missing']}"
        cells[5].text = ""
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    path = os.path.join(HERE, "PVSB_主场景回退抽检表.docx")
    doc.save(path)
    print("saved", path)


if __name__ == "__main__":
    main()
